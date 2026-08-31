#!/usr/bin/env python3
"""Render a .pptx back to PNGs so the exported deck can be checked by eye.

    python docs/portfolio/preview-pptx.py [deck.pptx] [outDir]

This is a verification tool, not a converter: it reads the shapes, fills, lines,
pictures and text runs that are actually stored in the file and draws them with
Pillow. PowerPoint's own text layout differs slightly (kerning, wrapping), so
treat the output as a geometry/colour check rather than a pixel reference.
Embedded SVG pictures are outlined and labelled instead of rasterised.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

SRC = Path(sys.argv[1] if len(sys.argv) > 1 else "docs/portfolio/portfolio-slides.pptx")
OUT = Path(sys.argv[2] if len(sys.argv) > 2 else "docs/portfolio/preview")
OUT.mkdir(parents=True, exist_ok=True)

W, H = 1920, 1080
EMU_PER_IN = 914400
SCALE = W / (13.333 * EMU_PER_IN)          # EMU -> px
PT2PX = 2.0                                 # 0.5pt per px on this stage

REG = "C:/Windows/Fonts/malgun.ttf"
BOLD = "C:/Windows/Fonts/malgunbd.ttf"
_font_cache: dict[tuple[str, int], ImageFont.FreeTypeFont] = {}


def font(size_px: int, bold: bool) -> ImageFont.FreeTypeFont:
    key = (BOLD if bold else REG, max(size_px, 6))
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype(key[0], key[1])
    return _font_cache[key]


def px(emu) -> float:
    return float(emu) * SCALE


NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def alpha_of(el, tag: str) -> float:
    """Read <a:alpha> inside a fill/line element; 1.0 when absent."""
    node = el.find(f".//{{{NS['a']}}}{tag}")
    if node is None:
        return 1.0
    a = node.find(f".//{{{NS['a']}}}alpha")
    return int(a.get("val")) / 100000 if a is not None else 1.0


def srgb(el, tag: str):
    node = el.find(f".//{{{NS['a']}}}{tag}")
    if node is None:
        return None
    c = node.find(f".//{{{NS['a']}}}srgbClr")
    if c is None:
        return None
    v = c.get("val")
    return tuple(int(v[i:i + 2], 16) for i in (0, 2, 4))


def slide_background(slide) -> tuple[int, int, int]:
    bg = slide._element.find(f".//{{{NS['p']}}}bg")
    if bg is None:
        return (247, 246, 243)
    stops = bg.findall(f".//{{{NS['a']}}}srgbClr")
    if not stops:
        return (247, 246, 243)
    v = stops[0].get("val")
    return tuple(int(v[i:i + 2], 16) for i in (0, 2, 4))


def draw_text(layer: Image.Image, shape) -> None:
    tf = shape.text_frame
    d = ImageDraw.Draw(layer)
    x0, y0 = px(shape.left), px(shape.top)
    box_w, box_h = px(shape.width), px(shape.height)

    lines: list[list[tuple[str, ImageFont.FreeTypeFont, tuple]]] = []
    line: list[tuple[str, ImageFont.FreeTypeFont, tuple]] = []
    width = 0.0
    line_px = 0.0

    for para in tf.paragraphs:
        for run in para.runs:
            size = run.font.size.pt * PT2PX if run.font.size else 20
            line_px = max(line_px, size * 1.35)
            f = font(int(size), bool(run.font.bold))
            col = (43, 45, 48)
            try:
                if run.font.color and run.font.color.rgb is not None:
                    col = tuple(run.font.color.rgb)
            except Exception:
                pass
            for word in run.text.split(" "):
                if not word:
                    continue
                w = d.textlength(word + " ", font=f)
                if width + w > box_w and line:
                    lines.append(line)
                    line, width = [], 0.0
                line.append((word + " ", f, col))
                width += w
            if run.text.endswith(" ") is False and run == para.runs[-1]:
                pass
        lines.append(line)
        line, width = [], 0.0

    lines = [ln for ln in lines if ln]
    total = len(lines) * line_px
    anchor = str(tf.vertical_anchor or "")
    y = y0 + (box_h - total) / 2 if "MIDDLE" in anchor else y0

    for ln in lines:
        lw = sum(d.textlength(t, font=f) for t, f, _ in ln)
        align = str(para.alignment or "")
        if "CENTER" in align:
            x = x0 + (box_w - lw) / 2
        elif "RIGHT" in align:
            x = x0 + box_w - lw
        else:
            x = x0
        for t, f, col in ln:
            d.text((x, y), t, font=f, fill=col)
            x += d.textlength(t, font=f)
        y += line_px


def render(slide, index: int) -> None:
    img = Image.new("RGB", (W, H), slide_background(slide))

    for shape in slide.shapes:
        x, y = px(shape.left), px(shape.top)
        w, h = px(shape.width), px(shape.height)
        box = [x, y, x + w, y + h]
        layer = Image.new("RGBA", (W, H), (0, 0, 0, 0))
        d = ImageDraw.Draw(layer)

        if shape.__class__.__name__ == "Picture":
            # reach the part directly — python-pptx tries to PIL-probe SVG blobs
            R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
            blip = shape._element.find(f".//{{{NS['a']}}}blip")
            blob = shape.part.related_part(blip.get(R)).blob if blip is not None else b""
            if blob[:5] == b"<?xml" or blob[:4] == b"<svg":
                d.rectangle(box, outline=(120, 140, 155, 255), width=3)
                d.text((x + 12, y + 12), "SVG (vector)", font=font(22, True), fill=(120, 140, 155, 255))
            else:
                pic = Image.open(io.BytesIO(blob)).convert("RGBA")
                pic = pic.resize((max(int(w), 1), max(int(h), 1)))
                layer.alpha_composite(pic, (int(x), int(y)))
            img = Image.alpha_composite(img.convert("RGBA"), layer).convert("RGB")
            continue

        sp = shape._element
        # fill and line live in <p:spPr> only — searching the whole shape would
        # pick up <a:solidFill> from the text runs and paint the box that colour
        spPr = sp.find(f".//{{{NS['p']}}}spPr")
        fill_el = spPr.find(f"{{{NS['a']}}}solidFill") if spPr is not None else None
        fill_c = srgb(fill_el, "solidFill") if fill_el is not None else None
        if fill_el is not None:
            c = fill_el.find(f"{{{NS['a']}}}srgbClr")
            if c is not None:
                v = c.get("val")
                fill_c = tuple(int(v[i:i + 2], 16) for i in (0, 2, 4))
        fill_a = alpha_of(fill_el, "srgbClr") if fill_el is not None else 0.0

        ln_el = spPr.find(f"{{{NS['a']}}}ln") if spPr is not None else None
        line_c = None
        line_a = 1.0
        if ln_el is not None:
            c = ln_el.find(f".//{{{NS['a']}}}srgbClr")
            if c is not None:
                v = c.get("val")
                line_c = tuple(int(v[i:i + 2], 16) for i in (0, 2, 4))
                a = c.find(f"{{{NS['a']}}}alpha")
                line_a = int(a.get("val")) / 100000 if a is not None else 1.0
            if ln_el.find(f"{{{NS['a']}}}noFill") is not None:
                line_c = None
        line_w = 1.0
        if ln_el is not None and ln_el.get("w"):
            line_w = int(ln_el.get("w")) / 12700 * PT2PX

        prst = sp.find(f".//{{{NS['a']}}}prstGeom")
        kind = prst.get("prst") if prst is not None else "rect"

        if fill_c and fill_a > 0.02:
            rgba = (*fill_c, int(255 * fill_a))
            if kind == "ellipse":
                d.ellipse(box, fill=rgba)
            elif kind == "roundRect":
                d.rounded_rectangle(box, radius=min(w, h) * 0.28, fill=rgba)
            else:
                d.rectangle(box, fill=rgba)
        if line_c:
            rgba = (*line_c, int(255 * line_a))
            if kind == "ellipse":
                d.ellipse(box, outline=rgba, width=max(int(line_w), 1))
            elif kind == "roundRect":
                d.rounded_rectangle(box, radius=min(w, h) * 0.28, outline=rgba, width=max(int(line_w), 1))
            else:
                d.rectangle(box, outline=rgba, width=max(int(line_w), 1))

        if shape.has_text_frame and shape.text_frame.text.strip():
            draw_text(layer, shape)

        img = Image.alpha_composite(img.convert("RGBA"), layer).convert("RGB")

    path = OUT / f"_ppt-{index:02d}.png"
    img.save(path)
    print(f"rendered {path}")


prs = Presentation(str(SRC))
for i, s in enumerate(prs.slides, start=1):
    render(s, i)
